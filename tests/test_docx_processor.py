import asyncio
from io import BytesIO
from pathlib import Path

import pytest

from redlines.utils.docx_processor import DOCXProcessor, DocxProcessingMethod
from redlines.utils.markdown_processor import MarkdownProcessor
from redlines.utils.styles import Styles

pytest.importorskip('docx')
from docx import Document as DocxDocument  # noqa: E402  (import after skip)


def test_manual_markdown_to_docx_conversion() -> None:
    processor = DOCXProcessor(
        styles=Styles(),
        markdown_processor=MarkdownProcessor(),
        html4docx_enabled=False,
        manual_enabled=True,
        preferred_method=DocxProcessingMethod.MANUAL,
    )
    docx_bytes = asyncio.run(processor.markdown_to_docx('# Title\n\nParagraph', method=DocxProcessingMethod.MANUAL))
    document = DocxDocument(BytesIO(docx_bytes))
    assert document.paragraphs[0].text == 'Title'
    assert 'Paragraph' in document.paragraphs[1].text


def test_docx_to_markdown_uses_sample_document_a() -> None:
    processor = DOCXProcessor(styles=Styles(), markdown_processor=MarkdownProcessor())
    sample_path = Path('tests/documents/PlainTextFile/forredlinetest_a.docx')
    payload = sample_path.read_bytes()
    markdown = asyncio.run(processor.docx_to_markdown(payload))
    assert 'Agreement Draft' in markdown
    assert 'respective obligations in good faith' in markdown


def test_docx_to_markdown_detects_delta_between_samples() -> None:
    processor = DOCXProcessor(styles=Styles(), markdown_processor=MarkdownProcessor())
    path_a = Path('tests/documents/PlainTextFile/forredlinetest_a.docx')
    path_b = Path('tests/documents/PlainTextFile/forredlinetest_b.docx')
    markdown_a = asyncio.run(processor.docx_to_markdown(path_a.read_bytes()))
    markdown_b = asyncio.run(processor.docx_to_markdown(path_b.read_bytes()))
    assert 'respective obligations in good faith' in markdown_a
    assert 'cooperate fully' in markdown_b
    assert markdown_a != markdown_b
