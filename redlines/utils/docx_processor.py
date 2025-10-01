from __future__ import annotations

import asyncio
import contextlib
import copy
import logging
import re
from enum import Enum
from io import BytesIO
from typing import Any, Optional

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from bs4 import BeautifulSoup, NavigableString
try:  # pragma: no cover - optional dependency
    from html4docx import HtmlToDocx
except ImportError:  # pragma: no cover - optional dependency
    HtmlToDocx = None

from .markdown_processor import MarkdownProcessor
from .styles import Styles
from .docx_track_changes import TrackChangesBuilder

logger = logging.getLogger(__name__)

DOCX_DEFAULT_TABLE_STYLE = 'Table Grid'


class DocxProcessingMethod(Enum):
    HTML4DOCX = 'html4docx'
    MANUAL = 'manual'


class DOCXProcessor:
    """High-fidelity DOCX conversion pipeline with async orchestration.

    The processor supports both html4docx-driven rendering and a resilient
    manual converter so that Markdown, HTML and text inputs can all be mapped to
    DOCX without pulling in Microsoft Office automation.
    """

    def __init__(
        self,
        *,
        styles: Optional[Styles] = None,
        markdown_processor: Optional[MarkdownProcessor] = None,
        html4docx_enabled: bool = True,
        manual_enabled: bool = True,
        preferred_method: DocxProcessingMethod = DocxProcessingMethod.HTML4DOCX,
        max_concurrent_operations: int = 6,
    ) -> None:
        self.styles = styles or Styles()
        self.markdown_processor = markdown_processor or MarkdownProcessor()
        self.html4docx_enabled = html4docx_enabled and HtmlToDocx is not None
        self.manual_enabled = manual_enabled
        self.preferred_method = preferred_method
        self._semaphore = asyncio.Semaphore(max_concurrent_operations)
        self._table_style = DOCX_DEFAULT_TABLE_STYLE

    # ------------------------------------------------------------------
    # Public conversion API
    # ------------------------------------------------------------------
    async def markdown_to_docx(
        self,
        markdown_content: str,
        *,
        method: Optional[DocxProcessingMethod] = None,
    ) -> bytes:
        """Render Markdown into DOCX, trying the configured backends in order."""
        async with self._semaphore:
            conversion_order = self._conversion_order(method)
            html = await self.markdown_processor.to_html(markdown_content, hard_wrap=True)
            for approach in conversion_order:
                try:
                    if approach is DocxProcessingMethod.HTML4DOCX and self.html4docx_enabled:
                        return await self._html_to_docx_html4docx(html)
                    if approach is DocxProcessingMethod.MANUAL and self.manual_enabled:
                        return await self._markdown_to_docx_manual(markdown_content)
                except Exception as exc:
                    logger.debug('%s conversion failed: %s', approach.value, exc)
                    continue
            msg = 'All DOCX conversion strategies failed'
            raise RuntimeError(msg)

    async def html_to_docx(
        self,
        html_content: str,
        *,
        method: Optional[DocxProcessingMethod] = None,
    ) -> bytes:
        """Convert HTML content to DOCX via html4docx or manual Markdown route."""
        async with self._semaphore:
            if self._contains_track_changes(html_content):
                try:
                    return await self._html_to_docx_track_changes(html_content)
                except Exception as exc:  # pragma: no cover - defensive
                    logger.debug('Track changes conversion failed, falling back: %s', exc)
            conversion_order = self._conversion_order(method)
            for approach in conversion_order:
                try:
                    if approach is DocxProcessingMethod.HTML4DOCX and self.html4docx_enabled:
                        return await self._html_to_docx_html4docx(html_content)
                    if approach is DocxProcessingMethod.MANUAL and self.manual_enabled:
                        markdown = await self.markdown_processor.from_html(html_content)
                        return await self._markdown_to_docx_manual(markdown)
                except Exception as exc:
                    logger.debug('%s conversion failed: %s', approach.value, exc)
                    continue
            msg = 'All DOCX conversion strategies failed'
            raise RuntimeError(msg)

    async def text_to_docx(
        self,
        text_content: str,
        *,
        method: Optional[DocxProcessingMethod] = None,
    ) -> bytes:
        """Convert plain text to DOCX after normalising to Markdown."""
        markdown = await self.markdown_processor.from_text(text_content)
        return await self.markdown_to_docx(markdown, method=method)

    async def docx_to_markdown(self, payload: bytes) -> str:
        """Extract Markdown from a DOCX payload using the shared processor."""
        return await self.markdown_processor.from_docx(payload)

    async def docx_to_text(self, payload: bytes) -> str:
        """Return plain text extracted from a DOCX payload."""
        markdown = await self.docx_to_markdown(payload)
        return await self.markdown_processor.to_text(markdown)

    async def docx_to_html(self, payload: bytes) -> str:
        """Return HTML extracted from a DOCX payload via Markdown rendering."""
        markdown = await self.docx_to_markdown(payload)
        return await self.markdown_processor.to_html(markdown)

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------
    def _conversion_order(self, preferred: Optional[DocxProcessingMethod]) -> list[DocxProcessingMethod]:
        if preferred:
            base = [preferred]
        else:
            base = [self.preferred_method]
        fallback = DocxProcessingMethod.MANUAL if base[0] is DocxProcessingMethod.HTML4DOCX else DocxProcessingMethod.HTML4DOCX
        return base + [fallback]

    @staticmethod
    def _contains_track_changes(html_content: str) -> bool:
        lowered = html_content.lower()
        return '<ins' in lowered or '<del' in lowered

    async def _html_to_docx_html4docx(self, html_content: str) -> bytes:
        if HtmlToDocx is None:  # pragma: no cover - defensive
            msg = 'html4docx backend is not available'
            raise RuntimeError(msg)
        parser = HtmlToDocx()
        with contextlib.suppress(Exception):
            parser.table_style = self._table_style
        document = await asyncio.to_thread(parser.parse_html_string, html_content)
        self.styles.configure_document(document)
        self._append_configured_text(document)
        buffer = BytesIO()
        await asyncio.to_thread(document.save, buffer)
        buffer.seek(0)
        return buffer.read()

    async def _markdown_to_docx_manual(self, markdown_text: str) -> bytes:
        """Manual Markdown-to-DOCX fallback that mirrors typical legal formatting."""
        doc = Document()
        self.styles.configure_document(doc)

        lines = markdown_text.split('\n')
        current_paragraph = None
        in_code_block = False
        code_block_content: list[str] = []
        table_data: list[list[str]] = []
        in_table = False

        for line in lines:
            stripped = line.strip()

            if stripped.startswith('```'):
                if in_code_block:
                    code_text = '\n'.join(code_block_content)
                    paragraph = doc.add_paragraph()
                    self._format_paragraph(paragraph)
                    run = paragraph.add_run(code_text)
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                    in_code_block = False
                    code_block_content = []
                else:
                    in_code_block = True
                current_paragraph = None
                continue

            if in_code_block:
                code_block_content.append(line)
                continue

            if not stripped:
                current_paragraph = None
                continue

            if stripped.startswith('|') and stripped.endswith('|'):
                if not in_table:
                    in_table = True
                    table_data = []
                cells = [cell.strip() for cell in stripped.strip('|').split('|')]
                if all(re.fullmatch(r'[:\-]+', cell) for cell in cells):
                    continue
                table_data.append(cells)
                continue
            if in_table:
                self._add_table(doc, table_data)
                table_data = []
                in_table = False

            if stripped.startswith('#'):
                level = len(stripped) - len(stripped.lstrip('#'))
                content = stripped[level:].strip()
                paragraph = doc.add_heading(content, level=min(level, 6))
                self.styles.apply_heading(paragraph, min(level, 6))
                current_paragraph = None
                continue

            # Ordered list
            if re.match(r'^\d+\.\s+', stripped):
                content = re.sub(r'^\d+\.\s+', '', stripped)
                paragraph = doc.add_paragraph(content, style='List Number')
                self._format_paragraph(paragraph)
                current_paragraph = None
                continue

            # Unordered list
            if stripped.startswith(('- ', '* ', '+ ')):
                content = re.sub(r'^[\-\*\+]\s+', '', stripped)
                paragraph = doc.add_paragraph(content, style='List Bullet')
                self._format_paragraph(paragraph)
                current_paragraph = None
                continue

            if stripped.startswith('>'):
                content = stripped.lstrip('>').strip()
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Pt(24)
                paragraph.paragraph_format.right_indent = Pt(24)
                self._format_paragraph(paragraph)
                self._add_inline_markdown(paragraph, content)
                current_paragraph = None
                continue

            if current_paragraph is None:
                current_paragraph = doc.add_paragraph()
                self._format_paragraph(current_paragraph)
            else:
                current_paragraph.add_run(' ')
            self._add_inline_markdown(current_paragraph, stripped)

        if in_table and table_data:
            self._add_table(doc, table_data)

        self._append_configured_text(doc)

        buffer = BytesIO()
        await asyncio.to_thread(doc.save, buffer)
        buffer.seek(0)
        return buffer.read()

    async def _html_to_docx_track_changes(self, html_content: str) -> bytes:
        """Generate DOCX output using native track changes elements."""
        soup = BeautifulSoup(html_content, 'html.parser')
        doc = Document()
        self.styles.configure_document(doc)
        builder = TrackChangesBuilder()

        body = soup.body or soup
        for element in body.children:
            if isinstance(element, NavigableString):
                text = str(element).strip()
                if not text:
                    continue
                paragraph = doc.add_paragraph(text)
                self._format_paragraph(paragraph)
                continue

            if not getattr(element, 'name', None):
                continue

            name = element.name.lower()
            if name in {'p', 'div'}:
                paragraph = doc.add_paragraph()
                self._format_paragraph(paragraph)
                self._append_nodes_to_paragraph(doc, paragraph, element, builder)
            elif name in {'ul', 'ol'}:
                style = 'List Number' if name == 'ol' else 'List Bullet'
                for li in element.find_all('li', recursive=False):
                    paragraph = doc.add_paragraph(style=style)
                    self._format_paragraph(paragraph)
                    self._append_nodes_to_paragraph(doc, paragraph, li, builder)
            elif name.startswith('h') and name[1:].isdigit():
                level = int(name[1:])
                paragraph = doc.add_paragraph()
                try:
                    self.styles.apply_heading(paragraph, level)
                except ValueError:
                    self._format_paragraph(paragraph)
                self._append_nodes_to_paragraph(doc, paragraph, element, builder)
            else:
                paragraph = doc.add_paragraph()
                self._format_paragraph(paragraph)
                self._append_nodes_to_paragraph(doc, paragraph, element, builder)

        buffer = BytesIO()
        await asyncio.to_thread(doc.save, buffer)
        buffer.seek(0)
        return buffer.read()

    def _append_nodes_to_paragraph(
        self,
        document: DocumentType,
        paragraph,
        element,
        builder: TrackChangesBuilder,
        formatting: Optional[Dict[str, object]] = None,
    ) -> None:
        formatting = formatting or {}

        for child in element.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text:
                    run = paragraph.add_run(text)
                    self._apply_run_formatting(run, formatting)
                continue

            if not getattr(child, 'name', None):
                continue

            name = child.name.lower()

            if name == 'br':
                paragraph.add_run().add_break()
                continue

            updated_formatting = self._derive_formatting(formatting, child)

            if name == 'ins':
                runs = self._collect_inline_runs(child, updated_formatting)
                if runs:
                    builder.append_insertion(paragraph, runs)
                continue

            if name == 'del':
                runs = self._collect_inline_runs(child, updated_formatting)
                if runs:
                    builder.append_deletion(paragraph, runs)
                continue

            if name in {'strong', 'b', 'em', 'i', 'u', 'span', 'sup', 'sub', 'code', 's'}:
                self._append_nodes_to_paragraph(document, paragraph, child, builder, updated_formatting)
                continue

            if name in {'p', 'div'}:
                new_paragraph = document.add_paragraph()
                self._format_paragraph(new_paragraph)
                self._append_nodes_to_paragraph(document, new_paragraph, child, builder)
                continue

            if name in {'ul', 'ol'}:
                style = 'List Number' if name == 'ol' else 'List Bullet'
                for li in child.find_all('li', recursive=False):
                    list_paragraph = document.add_paragraph(style=style)
                    self._format_paragraph(list_paragraph)
                    self._append_nodes_to_paragraph(document, list_paragraph, li, builder)
                continue

            if name.startswith('h') and name[1:].isdigit():
                level = int(name[1:])
                heading_paragraph = document.add_paragraph()
                try:
                    self.styles.apply_heading(heading_paragraph, level)
                except ValueError:
                    self._format_paragraph(heading_paragraph)
                self._append_nodes_to_paragraph(document, heading_paragraph, child, builder)
                continue

            self._append_nodes_to_paragraph(document, paragraph, child, builder, updated_formatting)

    def _collect_inline_runs(
        self,
        element,
        formatting: Dict[str, object],
    ) -> List[Dict[str, object]]:
        runs: List[Dict[str, object]] = []
        for child in element.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text:
                    runs.append(self._build_run_dict(text, formatting))
                continue

            if not getattr(child, 'name', None):
                continue

            name = child.name.lower()
            if name == 'br':
                runs.append({'break': True})
                continue

            nested_formatting = self._derive_formatting(formatting, child)

            if name in {'ins', 'del'}:
                runs.extend(self._collect_inline_runs(child, nested_formatting))
                continue

            runs.extend(self._collect_inline_runs(child, nested_formatting))

        return runs

    @staticmethod
    def _build_run_dict(text: str, formatting: Dict[str, object]) -> Dict[str, object]:
        run: Dict[str, object] = {'text': text}
        run.update({k: v for k, v in formatting.items() if v})
        return run

    @staticmethod
    def _derive_formatting(
        base: Dict[str, object], element
    ) -> Dict[str, object]:
        fmt = dict(base)
        name = element.name.lower()

        if name in {'strong', 'b'}:
            fmt['bold'] = True
        if name in {'em', 'i'}:
            fmt['italic'] = True
        if name == 'u':
            fmt['underline'] = True
        if name in {'s', 'strike'}:
            fmt['strike'] = True
        if name == 'code':
            fmt['font'] = 'Courier New'
            fmt['size'] = 10

        style_attr = element.get('style')
        if style_attr:
            for rule in style_attr.split(';'):
                if ':' not in rule:
                    continue
                key, value = rule.split(':', 1)
                key = key.strip().lower()
                value = value.strip().lower()
                if key == 'font-weight' and 'bold' in value:
                    fmt['bold'] = True
                elif key == 'font-style' and 'italic' in value:
                    fmt['italic'] = True
                elif key == 'text-decoration':
                    if 'underline' in value:
                        fmt['underline'] = True
                    if 'line-through' in value:
                        fmt['strike'] = True
                elif key == 'font-family':
                    fmt['font'] = value.strip("'\"")
                elif key == 'font-size':
                    size_value = value.replace('pt', '').strip()
                    try:
                        fmt['size'] = float(size_value)
                    except ValueError:
                        pass
                elif key == 'color':
                    fmt['color'] = value.lstrip('#')

        return fmt

    @staticmethod
    def _apply_run_formatting(run, formatting: Dict[str, object]) -> None:
        if formatting.get('bold'):
            run.bold = True
        if formatting.get('italic'):
            run.italic = True
        if formatting.get('underline'):
            run.underline = True
        if formatting.get('strike'):
            run.font.strike = True
        if font := formatting.get('font'):
            run.font.name = str(font)
        if size := formatting.get('size'):
            try:
                run.font.size = Pt(float(size))
            except (TypeError, ValueError):
                pass
        if color := formatting.get('color'):
            try:
                hex_color = str(color).lstrip('#')
                if len(hex_color) == 6:
                    run.font.color.rgb = RGBColor.from_string(hex_color.upper())
            except Exception:
                pass

    def _add_table(self, doc: DocumentType, data: list[list[str]]) -> None:
        if not data:
            return
        rows = len(data)
        cols = len(data[0])
        table_obj = doc.add_table(rows=rows, cols=cols)
        table_obj.style = self._table_style

        for r_idx, row_data in enumerate(data):
            for c_idx, cell_text in enumerate(row_data):
                cell = table_obj.cell(r_idx, c_idx)
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    self._format_paragraph(paragraph)
                    for run in paragraph.runs:
                        run.font.bold = r_idx == 0

    def _add_inline_markdown(self, paragraph, text: str) -> None:
        pattern = r'(\*\*(.+?)\*\*)|(\*(.+?)\*)|(__(.+?)__)|(\_(.+?)\_)|(`(.+?)`)'  # bold, italic, code
        position = 0
        for match in re.finditer(pattern, text):
            if match.start() > position:
                self._add_plain_run(paragraph, text[position : match.start()])
            if match.group(1):
                self._add_bold_run(paragraph, match.group(2))
            elif match.group(3):
                self._add_italic_run(paragraph, match.group(4))
            elif match.group(5):
                self._add_bold_run(paragraph, match.group(6))
            elif match.group(7):
                self._add_italic_run(paragraph, match.group(8))
            elif match.group(9):
                run = paragraph.add_run(match.group(10))
                run.font.name = 'Courier New'
                run.font.size = Pt(10)
            position = match.end()
        if position < len(text):
            self._add_plain_run(paragraph, text[position:])

    def _add_plain_run(self, paragraph, content: str) -> None:
        run = paragraph.add_run(content)
        self._format_run(run)

    def _add_bold_run(self, paragraph, content: str) -> None:
        run = paragraph.add_run(content)
        self._format_run(run, {'bold': True})

    def _add_italic_run(self, paragraph, content: str) -> None:
        run = paragraph.add_run(content)
        self._format_run(run, {'italic': True})

    def _format_paragraph(self, paragraph) -> None:
        self.styles.apply_paragraph_defaults(paragraph)

    def _append_configured_text(self, document: DocumentType) -> None:
        append_text = self.styles.config.get('append_text') or ''
        if not append_text.strip():
            return
        for line in append_text.split('\n'):
            paragraph = document.add_paragraph(line)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._format_paragraph(paragraph)

    def _format_run(self, run, overrides: Optional[dict[str, Any]] = None) -> None:
        font_cfg = copy.deepcopy(self.styles.config.get('font', {}))
        if overrides:
            Styles._deep_update(font_cfg, overrides)
        Styles._apply_run_font(run.font, font_cfg)


__all__ = ['DOCXProcessor', 'DocxProcessingMethod']
